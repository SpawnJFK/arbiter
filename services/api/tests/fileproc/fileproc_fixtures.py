"""Programmatic fixtures for the file format tests. No binaries are checked in.

DOCX and XLSX are written as raw XML in a zip so the tests control exactly which
constructs appear (split runs, fields, footnotes, tracked changes). PPTX comes from
python-pptx so it carries the full layout/master structure of a real deck.
"""

from __future__ import annotations

import base64
import io
import json
import zipfile

from arbiter.fileproc.base import Content, ExtractionResult, InlineCode, TargetMap, join_unit, to_tagged

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = (
    f'xmlns:w="{W}" xmlns:r="{R}" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _zip(parts: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data.encode() if isinstance(data, str) else data)
    return buf.getvalue()


def _p(inner: str, ppr: str = "") -> str:
    return f"<w:p>{ppr}{inner}</w:p>"


def _r(text: str, rpr: str = "", rsid: str = "") -> str:
    attr = f' w:rsidR="{rsid}"' if rsid else ""
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return f"<w:r{attr}>{f'<w:rPr>{rpr}</w:rPr>' if rpr else ''}<w:t{space}>{text}</w:t></w:r>"


DRAWING = (
    '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="9525" cy="9525"/>'
    '<wp:docPr id="1" name="Picture 1"/><a:graphic><a:graphicData '
    'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr>'
    '<pic:cNvPr id="0" name="logo.png"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
    '<a:blip r:embed="rId20"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr>'
    '<a:xfrm><a:off x="0" y="0"/><a:ext cx="9525" cy="9525"/></a:xfrm><a:prstGeom prst="rect"/>'
    "</pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing>"
)


def docx_body(tracked: bool = False, extra: str = "") -> str:
    paras = [
        # split runs with identical formatting but different rsids, plus proofing noise
        _p(
            _r("This is ", "", "00A1")
            + "<w:proofErr w:type=\"spellStart\"/>"
            + _r("im", "<w:b/>", "00A2")
            + _r("portant", "<w:b/><w:lang w:val=\"en-GB\"/>", "00A3")
            + "<w:proofErr w:type=\"spellEnd\"/>"
            + _r(" and ", "", "00A4")
            + _r("styled", "<w:i/>")
            + _r(" text in one sentence. It has a second sentence.")
        ),
        _p(
            _r("Visit ")
            + '<w:hyperlink r:id="rId10" w:history="1">'
            + _r("our website", '<w:rStyle w:val="Hyperlink"/>')
            + "</w:hyperlink>"
            + _r(" for more details.")
        ),
        _p(
            _r("Name:")
            + "<w:r><w:tab/></w:r>"
            + _r("John Smith")
            + '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r>'
        ),
        _p(
            _r("Page ")
            + '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            + '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
            + '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            + _r("1")
            + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
            + _r(" of the annual report")
        ),
        _p(_r("Logo here: ") + f"<w:r>{DRAWING}</w:r>" + _r(" and the text goes on.")),
        _p(
            '<w:bookmarkStart w:id="0" w:name="intro"/>'
            + _r("Bookmarked text")
            + '<w:bookmarkEnd w:id="0"/>'
        ),
        _p(_r("Line one") + "<w:r><w:br/></w:r>" + _r("Line two")),
        _p(_r("12345")),
        _p(_r("   ")),
        "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid><w:tr>"
        f"<w:tc>{_p(_r('First cell'))}</w:tc>"
        f"<w:tc>{_p(_r('Second cell'))}"
        "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr>"
        f"<w:tc>{_p(_r('Nested cell'))}</w:tc></w:tr></w:tbl>{_p('')}</w:tc>"
        "</w:tr></w:tbl>",
    ]
    if tracked:
        paras.append(_p(_r("Old text ") + f"<w:ins w:id=\"9\" w:author=\"x\">{_r('inserted')}</w:ins>"))
    body = "".join(paras) + extra
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:document {NS}><w:body>{body}'
        '<w:sectPr><w:headerReference w:type="default" r:id="rId1"/>'
        '<w:footerReference w:type="default" r:id="rId2"/></w:sectPr></w:body></w:document>'
    )


def make_docx(tracked: bool = False, extra: str = "") -> bytes:
    ct = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/header1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
        '<Override PartName="/word/footer1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/word/comments.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'officeDocument" Target="word/document.xml"/></Relationships>'
    )
    base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{base}header" Target="header1.xml"/>'
        f'<Relationship Id="rId2" Type="{base}footer" Target="footer1.xml"/>'
        f'<Relationship Id="rId3" Type="{base}footnotes" Target="footnotes.xml"/>'
        f'<Relationship Id="rId4" Type="{base}comments" Target="comments.xml"/>'
        f'<Relationship Id="rId10" Type="{base}hyperlink" Target="https://example.com" TargetMode="External"/>'
        f'<Relationship Id="rId20" Type="{base}image" Target="media/image1.png"/>'
        "</Relationships>"
    )
    header = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:hdr {NS}>{_p(_r("Company header"))}</w:hdr>'
    footer = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:ftr {NS}>'
        + _p(_r("Footer text, page ") + '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
        + "</w:ftr>"
    )
    footnotes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {NS}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
        + _r(" This is a footnote.")
        + "</w:p></w:footnote></w:footnotes>"
    )
    comments = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:comments {NS}>'
        f'<w:comment w:id="0" w:author="A">{_p(_r("Reviewer comment, not translated"))}</w:comment></w:comments>'
    )
    return _zip({
        "[Content_Types].xml": ct,
        "_rels/.rels": rels,
        "word/document.xml": docx_body(tracked, extra),
        "word/_rels/document.xml.rels": doc_rels,
        "word/header1.xml": header,
        "word/footer1.xml": footer,
        "word/footnotes.xml": footnotes,
        "word/comments.xml": comments,
        "word/media/image1.png": PNG,
    })


# --------------------------------------------------------------------------- XLSX

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


# A table of contents: one complex field spanning three paragraphs, with a nested PAGEREF field
# inside a hyperlink, followed by a paragraph holding a text box (with a VML fallback copy).
TOC_AND_TEXTBOX = (
    _p(_r("Contents ") + '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
       '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" </w:instrText></w:r>'
       '<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
    + _p('<w:hyperlink w:anchor="_Toc1">' + _r("Introduction chapter")
         + '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGEREF _Toc1 </w:instrText></w:r>'
         '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _r("2")
         + '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink>')
    + _p('<w:r><w:fldChar w:fldCharType="end"/></w:r>' + _r("After the table of contents."))
    + _p(_r("Box: ") + "<w:r><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><w:txbxContent>"
         + _p(_r("Text box words")) + "</w:txbxContent></w:drawing></mc:Choice><mc:Fallback><w:pict><w:txbxContent>"
         + _p(_r("Text box words")) + "</w:txbxContent></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
         + _r(" after the box."))
)


def make_xlsx() -> bytes:
    shared = [
        "<si><t>Product name</t></si>",
        "<si><t>Price</t></si>",
        '<si><r><t xml:space="preserve">Very </t></r><r><rPr><b/><sz val="11"/></rPr><t>strong</t></r>'
        '<r><t xml:space="preserve"> coffee beans</t></r></si>',
        "<si><t>12345</t></si>",
        "<si><t>Orphan string never shown</t></si>",
        "<si><t>Line_x000D_\nbreak cell</t></si>",
    ]
    sst = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="{S}" count="6" uniqueCount="6">'
        + "".join(shared)
        + "</sst>"
    )
    sheet = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="{S}"><sheetData>'
        '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
        '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2"><v>4.5</v></c></row>'
        '<row r="3"><c r="A3" t="s"><v>3</v></c><c r="B3"><f>SUM(B2:B2)</f><v>4.5</v></c>'
        '<c r="C3" t="str"><f>"Formula "&amp;"text"</f><v>Formula text</v></c></row>'
        '<row r="4"><c r="A4" t="inlineStr"><is><t>Inline greeting</t></is></c>'
        '<c r="B4" t="s"><v>5</v></c><c r="C4" t="s"><v>0</v></c></row>'
        "</sheetData></worksheet>"
    )
    workbook = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook xmlns="{S}" xmlns:r="{R}">'
        '<sheets><sheet name="Menu" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{base}worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{base}sharedStrings" Target="sharedStrings.xml"/>'
        "</Relationships>"
    )
    ct = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{base}officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    return _zip({
        "[Content_Types].xml": ct,
        "_rels/.rels": rels,
        "xl/workbook.xml": workbook,
        "xl/_rels/workbook.xml.rels": wb_rels,
        "xl/worksheets/sheet1.xml": sheet,
        "xl/sharedStrings.xml": sst,
    })


# --------------------------------------------------------------------------- PPTX


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly results"
    body = slide.placeholders[1].text_frame
    p = body.paragraphs[0]
    p.add_run().text = "Revenue grew "
    r = p.add_run()
    r.text = "strongly"
    r.font.bold = True
    p.add_run().text = " this quarter. Costs fell."
    p2 = body.add_paragraph()
    p2.add_run().text = "Read the "
    link = p2.add_run()
    link.text = "full report"
    link.hyperlink.address = "https://example.com/report"
    p2.add_run().text = " online"
    p2.add_line_break()
    p2.add_run().text = "Second line"
    body.add_paragraph().add_run().text = "2024"
    slide.notes_slide.notes_text_frame.text = "Speaker notes for the first slide."

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Team table"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Ana"
    table.cell(1, 1).text = "Lead engineer"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- text formats

HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Arbiter test page</title>
<meta name="description" content="A page used to test HTML extraction">
<style>p { color: red; }</style>
</head>
<body>
<h1>Welcome &amp; hello</h1>
<p>This is <b>bold</b> and <a href="/docs" class="x">a link</a>.<br>Next line with <img src="a.png" alt="A cat"> picture. Another sentence!</p>
<div>Intro text <p>Nested paragraph</p> tail text</div>
<ul><li>First <em>item</em></li><li>Second item</li></ul>
<pre>preformatted, not translated</pre>
<p translate="no">Brand name stays</p>
<p class="notranslate">Also stays</p>
<p>Code like <code>print()</code> and <span class="notranslate">ACME</span> inline.</p>
<script>var greeting = "do not translate";</script>
<form><input type="text" placeholder="Your name"><input type="submit" value="Send"></form>
<p title="Tooltip text">With title</p>
</body>
</html>
"""

MARKDOWN = """---
title: Front matter is not translated
---

# Main *title* here

Some **bold** text with `code` and a [link **inside**](http://x.com/a_(b)) here.
Continued line. Second sentence? Yes!

- item one with ![alt text](img.png)
- [ ] task item <b>html</b> and {{ var }}
  continuation

> quote line one
> - quoted list

| Col A | Col B |
|------|:-----:|
| cell 1 | `x` |
| snake_case_name | 42 |

```python
print("no")
```

    indented code

1. Numbered 5 * 3 = 15 see https://example.com/page.

Trailing paragraph without newline"""

JSON_DOC = {
    "app": {
        "title": "Hello {name}!",
        "items": "{count, plural, =0 {No items} one {# item} other {# items for {user}}}",
        "gender": "{g, select, female {She liked it} male {He liked it} other {They liked it}}",
        "fmt": "Downloaded %s of %1$d files",
        "tpl": "Hi {{user}}, you have ${count} messages",
        "url": "https://example.com",
        "number": "123",
        "list": ["First entry", 5, None, True, "Second entry"],
        "a/b~c": "Escaped key",
    },
    "empty": "",
}


def json_bytes(indent: int | str | None) -> bytes:
    return (json.dumps(JSON_DOC, indent=indent, ensure_ascii=False) + "\n").encode()


PO = r'''# Translation template for Arbiter.
msgid ""
msgstr ""
"Project-Id-Version: demo 1.0\n"
"Language: \n"
"MIME-Version: 1.0\n"
"Content-Type: text/plain; charset=UTF-8\n"
"Plural-Forms: nplurals=INTEGER; plural=EXPRESSION;\n"

# Shown on the start page
#. Developer comment
#: src/app.py:10
#, fuzzy, python-format
#| msgid "Old hello"
msgctxt "greeting"
msgid "Hello %(name)s, welcome!\n"
msgstr "Zdravo"

msgid ""
"Long text that "
"wraps over lines. Second {count}."
msgstr ""

#, c-format
msgid "One file"
msgid_plural "%d files"
msgstr[0] ""
msgstr[1] ""

msgid "1.0"
msgstr ""

#~ msgid "obsolete"
#~ msgstr "zastarelo"
'''

PO_TRANSLATED = r'''msgid ""
msgstr ""
"Language: en\n"
"Content-Type: text/plain; charset=UTF-8\n"
"Plural-Forms: nplurals=2; plural=(n != 1);\n"

msgid "Save file"
msgstr "Save file"

msgid "One file"
msgid_plural "%d files"
msgstr[0] "One file"
msgstr[1] "%d files"
'''

TXT = (
    "First paragraph has two sentences. This is the second one.\r\n"
    "It continues here.\r\n"
    "\r\n"
    "   Indented second paragraph!\r\n"
    "\r\n\r\n"
    "12345\r\n"
    "Last line without newline"
)

CSV = 'id;name;description\r\n1;Coffee;"Strong; dark roast"\r\n2;"Tea";Green tea, mild\r\n3;42;"He said ""hi"""\r\n'

XLIFF12 = """<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="app.properties" source-language="en" datatype="plaintext">
    <body>
      <trans-unit id="t1" resname="greeting">
        <source>Hello <g id="1">bold</g> world<x id="2"/>. Second sentence here.</source>
        <note>Shown on login</note>
      </trans-unit>
      <trans-unit id="t2">
        <source>Click <bx id="3" rid="1"/>here<ex id="4" rid="1"/> now.</source>
        <seg-source><mrk mtype="seg" mid="1">Click <bx id="3" rid="1"/>here<ex id="4" rid="1"/> now.</mrk></seg-source>
      </trans-unit>
      <trans-unit id="t3">
        <source>First part. Second part.</source>
        <seg-source><mrk mtype="seg" mid="1">First part.</mrk> <mrk mtype="seg" mid="2">Second part.</mrk></seg-source>
        <target>old</target>
      </trans-unit>
      <trans-unit id="t4" translate="no"><source>Do not touch</source></trans-unit>
      <trans-unit id="t5"><source>Value <ph id="5">%s</ph> left</source></trans-unit>
    </body>
  </file>
</xliff>
"""

XLIFF20 = """<?xml version="1.0" encoding="UTF-8"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:2.0" version="2.0" srcLang="en">
  <file id="f1">
    <unit id="u1" name="welcome">
      <notes><note>Greeting</note></notes>
      <segment><source>Hello <pc id="1">bold</pc> and <ph id="2"/> end.</source></segment>
      <ignorable><source> </source></ignorable>
      <segment state="initial"><source>Second <sc id="3"/>marked<ec startRef="3"/> segment.</source></segment>
    </unit>
    <unit id="u2" translate="no"><segment><source>Skip me</source></segment></unit>
    <unit id="u3"><segment><source>Already <mrk id="m1" type="term">done</mrk>.</source><target>Fertig.</target></segment></unit>
  </file>
</xliff>
"""


# --------------------------------------------------------------------------- helpers


def tagged(result: ExtractionResult) -> list[tuple[str, list[str], str]]:
    return [(u.unit_id, [to_tagged(s.content) for s in u.segments], u.context) for u in result.units]


def pseudo(content: Content, left: str = "[[", right: str = "]]") -> Content:
    """Pseudo-translate: wrap every non-blank text run, keep codes where they are."""
    return [f"{left}{r.upper()}{right}" if isinstance(r, str) and r.strip() else r for r in content]


def pseudo_targets(result: ExtractionResult, left: str = "[[", right: str = "]]") -> TargetMap:
    return {u.unit_id: [pseudo(s.content, left, right) for s in u.segments] for u in result.units}


def joined(unit) -> Content:  # noqa: ANN001
    return join_unit([s.content for s in unit.segments], unit.segments, unit.leading_ws)


def plain(content: Content) -> str:
    return "".join(r for r in content if isinstance(r, str))


def code_kinds(content: Content) -> list[tuple[str, str]]:
    return sorted((c.kind, c.display) for c in content if isinstance(c, InlineCode))
