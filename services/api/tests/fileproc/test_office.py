from __future__ import annotations

import io
import zipfile

import openpyxl
import pptx
from fileproc_fixtures import joined, make_pptx, make_xlsx, plain, tagged
from roundtrip import identity, translation

from arbiter.fileproc.pptx import PptxHandler
from arbiter.fileproc.xlsx import XlsxHandler


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


# --------------------------------------------------------------------------- XLSX


def test_xlsx_extraction() -> None:
    result = XlsxHandler().extract(make_xlsx(), "en")
    units = {uid: (segs, ctx) for uid, segs, ctx in tagged(result)}
    assert units["xl/sharedStrings.xml#0"] == (["Product name"], "sheet:Menu!A1")
    assert units["xl/sharedStrings.xml#2"][0] == ["Very ⟦1⟧strong⟦/1⟧ coffee beans"]
    assert units["xl/sharedStrings.xml#5"][0] == ["Line\r\nbreak cell"]
    assert units["xl/worksheets/sheet1.xml!A4"] == (["Inline greeting"], "sheet:Menu!A4")
    texts = [plain(joined(u)) for u in result.units]
    assert "12345" not in texts  # numeric only
    assert "Orphan string never shown" not in texts  # not referenced by any cell
    assert "Formula text" not in texts  # formula result
    assert all(len(u.segments) == 1 for u in result.units)


def test_xlsx_identity() -> None:
    data = make_xlsx()
    merged, _ = identity(XlsxHandler(), data)
    assert merged == data
    forced, _ = identity(XlsxHandler(rewrite_unchanged=True), data)
    before, after = _parts(data), _parts(forced)
    for name in before:
        if name not in ("xl/sharedStrings.xml", "xl/worksheets/sheet1.xml"):
            assert before[name] == after[name], name
    wb = openpyxl.load_workbook(io.BytesIO(forced), rich_text=True)
    assert str(wb["Menu"]["A4"].value) == "Inline greeting"


def test_xlsx_translation() -> None:
    data = make_xlsx()
    merged, _, _ = translation(XlsxHandler(), data)
    sst = _parts(merged)["xl/sharedStrings.xml"].decode()
    assert "<b/>" in sst and "[[STRONG]]" in sst
    assert "_x000D_" in sst  # carriage return re-escaped the Excel way
    assert "Orphan string never shown" in sst  # untouched
    wb = openpyxl.load_workbook(io.BytesIO(merged))
    ws = wb["Menu"]
    assert ws["A1"].value == "[[PRODUCT NAME]]"
    assert ws["C4"].value == "[[PRODUCT NAME]]"  # same shared string, translated once
    assert ws["A4"].value == "[[INLINE GREETING]]"
    assert ws["B3"].value == "=SUM(B2:B2)"
    assert ws["A3"].value == "12345"


# --------------------------------------------------------------------------- PPTX


def test_pptx_extraction() -> None:
    result = PptxHandler().extract(make_pptx(), "en")
    by_text = {plain(joined(u)): u for u in result.units}
    assert "Quarterly results" in by_text
    body = next(u for t, u in by_text.items() if t.startswith("Revenue"))
    assert [s for s in tagged_unit(body)] == ["Revenue grew ⟦1⟧strongly⟦/1⟧ this quarter.", "Costs fell."]
    link = next(u for t, u in by_text.items() if t.startswith("Read the"))
    assert tagged_unit(link) == ["Read the ⟦1⟧full report⟦/1⟧ online⟦2/⟧Second line"]
    assert link.segments[0].content[1].display.startswith("link")
    assert by_text["Speaker notes for the first slide."].context == "notes"
    assert by_text["Lead engineer"].context == "table"
    assert "2024" not in by_text


def tagged_unit(unit) -> list[str]:  # noqa: ANN001
    from arbiter.fileproc.base import to_tagged

    return [to_tagged(s.content) for s in unit.segments]


def test_pptx_identity() -> None:
    data = make_pptx()
    merged, _ = identity(PptxHandler(), data)
    assert merged == data
    forced, _ = identity(PptxHandler(rewrite_unchanged=True), data)
    before, after = _parts(data), _parts(forced)
    changed = {n for n in before if before[n] != after[n]}
    assert all("/slides/" in n or "/notesSlides/" in n for n in changed)
    prs = pptx.Presentation(io.BytesIO(forced))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    orig = [
        sh.text_frame.text
        for s in pptx.Presentation(io.BytesIO(data)).slides
        for sh in s.shapes
        if sh.has_text_frame
    ]
    assert texts == orig


def test_pptx_translation() -> None:
    data = make_pptx()
    merged, _, _ = translation(PptxHandler(), data)
    prs = pptx.Presentation(io.BytesIO(merged))
    slide = prs.slides[0]
    runs = [r for p in slide.placeholders[1].text_frame.paragraphs for r in p.runs]
    bold = [r for r in runs if "STRONGLY" in r.text]
    assert bold and bold[0].font.bold
    linked = [r for r in runs if "FULL REPORT" in r.text]
    assert linked and linked[0].hyperlink.address == "https://example.com/report"
    assert "[[SPEAKER NOTES" in slide.notes_slide.notes_text_frame.text
    table = next(sh for sh in prs.slides[1].shapes if sh.has_table).table
    assert table.cell(1, 1).text == "[[LEAD ENGINEER]]"
