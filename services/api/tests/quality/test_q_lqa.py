"""LQA scorecards: formula, job scorecard from reviewer + senate errors, exports, org settings."""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from sqlalchemy import select

from arbiter.api.app import create_app
from arbiter.models import Job, Organization, ReviewTask, Segment, SenateRun
from arbiter.pipeline.worker import run_until_idle
from arbiter.quality import scorecard

DOC = (
    b"# Release notes\n\n"
    b"The contract must be signed before the update. Open Settings and click Save.\n\n"
    b"Arbiter supports 12 file formats. Prices start at 3.5 EUR.\n"
)


def _setup(db):
    c = TestClient(create_app())
    r = c.post(
        "/v1/auth/register",
        json={"org_name": "Acme", "name": "PM", "email": "pm@example.com", "password": "secret-123"},
    )
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    f = c.post("/v1/files", headers=h, files={"file": ("n.md", DOC)}, data={"source_lang": "en"}).json()
    q = c.post("/v1/quotes", headers=h, json={"file_id": f["id"], "target_langs": ["sr"]}).json()
    job_id = c.post(
        "/v1/projects", headers=h, json={"name": "P", "quote_id": q["id"], "tier": "full"}
    ).json()["jobs"][0]["id"]
    run_until_idle()
    return c, h, db.get(Job, job_id)


def test_scorecard_formula_weights_and_threshold():
    rows = [
        {"dimension": "accuracy", "severity": "major"},
        {"dimension": "terminology", "severity": "minor"},
        {"dimension": "style", "severity": "neutral"},
    ]
    card = scorecard.scorecard(rows, 200, scorecard.default_settings())
    assert card["penalty"] == 6 and card["score"] == 97.0 and card["passed"]
    assert card["by_dimension"]["accuracy"]["major"] == 1 and card["by_severity"]["minor"] == 1
    heavy = {"threshold": 98.0, "weights": {"accuracy": 2.0}}
    card = scorecard.scorecard(rows, 200, heavy)
    assert card["penalty"] == 11 and card["score"] == 94.5 and not card["passed"]
    assert scorecard.scorecard([], 0, scorecard.default_settings())["score"] == 100.0
    assert scorecard.normalize_error({"dimension": "Fluency/Grammar", "severity": "minor"}, "x")[
        "dimension"
    ] == ("linguistic_conventions")
    assert (
        scorecard.normalize_error({"dimension": "whatever", "severity": "major"}, "x")["dimension"] == "other"
    )
    assert scorecard.normalize_error({"dimension": "accuracy", "severity": "huge"}, "x") is None


def test_job_lqa_from_reviewer_and_senate_errors_and_exports(db):
    c, h, job = _setup(db)
    segs = db.execute(select(Segment).where(Segment.job_id == job.id).order_by(Segment.seq)).scalars().all()
    tasks = {
        t.segment_id: t for t in db.execute(select(ReviewTask).where(ReviewTask.job_id == job.id)).scalars()
    }
    t1 = tasks[segs[1].id]
    t1.state, t1.errors = (
        "accepted",
        [
            {"dimension": "accuracy", "severity": "major", "span": "potpisan", "explanation": "wrong verb"},
            {"dimension": "fluency", "severity": "minor", "span": "", "explanation": "word order"},
        ],
    )
    t2 = tasks[segs[2].id]
    t2.state, t2.errors = "rejected", [{"dimension": "accuracy", "severity": "critical", "span": "x"}]
    db.add(
        SenateRun(
            segment_id=segs[1].id,
            job_id=job.id,
            purpose="review",
            outcome="errors_confirmed",
            confirmed=[
                {"dimension": "accuracy", "severity": "minor", "span": "Potpisan", "explanation": "dup"},
                {"dimension": "terminology", "severity": "major", "span": "ugovor", "explanation": "term"},
            ],
        )
    )
    db.commit()
    r = c.get(f"/v1/jobs/{job.id}/lqa", headers=h)
    assert r.status_code == 200, r.text
    card = r.json()
    # accuracy major (reviewer + senate dup counted once, higher severity), fluency minor, terminology major
    assert card["error_count"] == 3 and card["penalty"] == 11
    assert card["by_dimension"]["accuracy"]["major"] == 1 and card["by_dimension"]["accuracy"]["minor"] == 0
    assert card["sources"] == {"reviewer": 2, "senate": 2}
    assert card["word_count"] == job.word_count and card["threshold"] == 95.0
    expected = round(max(0.0, 100 - 11 * 100 / job.word_count), 2)
    assert card["score"] == expected and card["passed"] is (expected >= 95.0)
    dup = next(e for e in card["errors"] if e["dimension"] == "accuracy")
    assert dup["source"] == "reviewer+senate" and dup["seq"] == 2

    pdf = c.get(f"/v1/jobs/{job.id}/lqa.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    x = c.get(f"/v1/jobs/{job.id}/lqa.xlsx", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Errors"] and wb["Errors"].max_row == 4


def test_org_lqa_settings_and_module_gate(db):
    c, h, job = _setup(db)
    s = c.get("/v1/lqa/settings", headers=h).json()
    assert s["threshold"] == 95.0 and "design_and_markup" in s["dimensions"]
    r = c.put("/v1/lqa/settings", headers=h, json={"threshold": 99, "weights": {"terminology": 2}})
    assert r.status_code == 200 and r.json()["threshold"] == 99 and r.json()["weights"]["terminology"] == 2
    assert c.get(f"/v1/jobs/{job.id}/lqa", headers=h).json()["threshold"] == 99
    assert c.put("/v1/lqa/settings", headers=h, json={"weights": {"vibes": 1}}).status_code == 422
    assert c.put("/v1/lqa/settings", headers=h, json={"threshold": 101}).status_code == 422
    assert c.put("/v1/lqa/settings", headers=h, json={"reset": True}).json()["threshold"] == 95.0
    org = db.execute(select(Organization)).scalar_one()
    org.settings = {**(org.settings or {}), "modules": ["tms"]}
    db.commit()
    r = c.get(f"/v1/jobs/{job.id}/lqa", headers=h)
    assert r.status_code == 404 and r.json()["error"]["code"] == "module_disabled"
    h2 = {"Authorization": h["Authorization"]}
    assert c.get(f"/v1/jobs/{job.id}", headers=h2).status_code == 200
