"""Auditor: bilingual readers, the queued audit pipeline, report, filters, exports, usage."""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from lxml import etree
from sqlalchemy import select

from arbiter.api.app import create_app
from arbiter.auditor import bilingual
from arbiter.models import Audit, Organization, UsageRecord, WorkItem
from arbiter.pipeline.worker import run_until_idle

XLF12 = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
 <file original="ui.txt" source-language="en" target-language="de" datatype="plaintext"><body>
  <trans-unit id="1"><source>Click <g id="b1">Save</g> to keep 3 files.</source>
   <target>Klicken Sie auf <g id="b1">Speichern</g>, um 4 Dateien zu behalten.</target></trans-unit>
  <trans-unit id="2"><source>Open the contract.</source><target>Offnen Sie den Vertrag.</target></trans-unit>
  <trans-unit id="3"><source>Open the contract.</source><target>Offnen Sie die Vereinbarung.</target></trans-unit>
  <trans-unit id="4"><source>Image <x id="img"/> here.</source><target>Bild hier.</target></trans-unit>
  <trans-unit id="5"><source>Not translated yet.</source></trans-unit>
 </body></file>
</xliff>"""

XLF2 = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="2.0" xmlns="urn:oasis:names:tc:xliff:document:2.0" srcLang="en" trgLang="de">
 <file id="f1"><unit id="u1">
  <segment id="s1"><source>Hello <pc id="1">world</pc>.</source><target>Hallo <pc id="1">Welt</pc>.</target></segment>
  <segment id="s2"><source>Visit https://example.com today.</source><target>Besuchen Sie heute example.com.</target></segment>
 </unit></file>
</xliff>"""

TMX = b"""<?xml version="1.0" encoding="UTF-8"?>
<tmx version="1.4"><header srclang="en" segtype="sentence" datatype="plaintext" adminlang="en" creationtool="t"
 creationtoolversion="1" o-tmf="t"/>
<body><tu tuid="a"><tuv xml:lang="en"><seg>Save the file.</seg></tuv><tuv xml:lang="de-DE"><seg>Datei speichern.</seg></tuv></tu>
<tu tuid="b"><tuv xml:lang="en"><seg>Close <bpt i="1">&lt;b&gt;</bpt>now<ept i="1">&lt;/b&gt;</ept>.</seg></tuv>
<tuv xml:lang="de"><seg>Jetzt schliessen.</seg></tuv></tu></body></tmx>"""

CSV = b"source;target\nSave the file.;Datei speichern.\nPrice: 10 EUR;Preis: 12 EUR\n=1+1;=1+1\n"


def _client(db):
    c = TestClient(create_app())
    r = c.post(
        "/v1/auth/register",
        json={"org_name": "Acme", "name": "PM", "email": "pm@example.com", "password": "secret-123"},
    )
    return c, {"Authorization": f"Bearer {r.json()['token']}"}


def test_readers_xliff12_xliff2_tmx_csv_xlsx():
    s = bilingual.parse_xliff(XLF12)
    assert len(s) == 5 and s[0].ref == "1|1|"
    assert s[0].source_tagged == "Click ⟦b1⟧Save⟦/b1⟧ to keep 3 files."
    assert s[0].target_tagged.startswith("Klicken Sie auf ⟦b1⟧Speichern⟦/b1⟧")
    assert s[3].source_tagged == "Image ⟦img/⟧ here." and s[4].target_tagged == ""
    s2 = bilingual.parse_xliff(XLF2)
    assert [x.ref for x in s2] == ["1|u1|s1", "1|u1|s2"] and s2[0].target_tagged == "Hallo ⟦1⟧Welt⟦/1⟧."
    t = bilingual.parse_tmx(TMX, "en", "de")
    assert [x.target_tagged for x in t] == ["Datei speichern.", "Jetzt schliessen."]
    assert t[1].source_tagged == "Close ⟦1⟧now⟦/1⟧."
    c = bilingual.parse_csv(CSV)
    assert len(c) == 3 and c[1].source_tagged == "Price: 10 EUR" and c[1].ref == "row:3"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Original", "Translation"])
    ws.append(["Hello", "Hallo"])
    buf = io.BytesIO()
    wb.save(buf)
    x = bilingual.parse_xlsx(buf.getvalue())
    assert [(r.source_tagged, r.target_tagged) for r in x] == [("Hello", "Hallo")]
    assert bilingual.detect_format("a.sdlxliff") == "xliff" and bilingual.detect_format("a.docx") is None


def test_audit_xliff_end_to_end_report_filters_exports_and_ai_usage(db):
    c, h = _client(db)
    g = c.post("/v1/glossaries", headers=h, json={"name": "UI"}).json()
    c.post(
        f"/v1/glossaries/{g['id']}/terms",
        headers=h,
        json={"source_lang": "en", "target_lang": "de", "source_term": "contract", "target_term": "Vertrag"},
    )
    r = c.post(
        "/v1/audits",
        headers=h,
        files={"file": ("ui.xlf", XLF12)},
        data={"source_lang": "en", "target_lang": "de", "glossary_id": g["id"], "ai_mode": "full"},
    )
    assert r.status_code == 201, r.text
    a = r.json()
    assert a["state"] == "queued" and a["input_format"] == "xliff" and a["progress"] == 0
    assert c.get(f"/v1/audits/{a['id']}/report.pdf", headers=h).status_code == 409
    run_until_idle()
    a = c.get(f"/v1/audits/{a['id']}", headers=h).json()
    assert a["state"] == "done", a
    assert a["progress"] == 1.0 and a["segments_total"] == 5 and a["segments_done"] == 5
    assert a["verdict"] in ("pass", "fail") and a["mqm_score"] is not None
    summ = a["summary"]
    codes = summ["issues_by_code"]
    assert codes["number_mismatch"] == 1 and codes["tag_missing"] == 1 and codes["empty_target"] == 1
    assert codes["term_missing"] == 1  # "Vereinbarung" instead of the mandatory "Vertrag"
    assert codes["inconsistent_translation"] == 2 and summ["consistency"]["inconsistent_translations"] == 1
    assert summ["ai"]["mode"] == "full" and summ["ai"]["checked"] >= 1
    assert summ["scorecard"]["error_count"] >= 6 and a["verdict"] == "fail"
    usage = db.execute(select(UsageRecord)).scalars().all()
    assert len(usage) == 1 and usage[0].unit == "ai_unit" and usage[0].job_id is None
    assert str(usage[0].quantity) == a["ai_units"].rstrip("0").rstrip(".") or usage[0].quantity > 0

    segs = c.get(f"/v1/audits/{a['id']}/segments", headers=h).json()
    assert segs["total"] == 5 and segs["items"][0]["issues"]
    major = c.get(f"/v1/audits/{a['id']}/segments", headers=h, params={"min_severity": "major"}).json()
    assert major["total"] >= 3 and all(s["max_severity"] in ("major", "critical") for s in major["items"])
    term = c.get(f"/v1/audits/{a['id']}/segments", headers=h, params={"dimension": "terminology"}).json()
    assert {s["seq"] for s in term["items"]} >= {2}
    code = c.get(f"/v1/audits/{a['id']}/segments", headers=h, params={"code": "tag_missing"}).json()
    assert [s["seq"] for s in code["items"]] == [3]
    assert c.get(f"/v1/audits/{a['id']}/segments?severity=loud", headers=h).status_code == 422
    assert c.get(f"/v1/audits/{a['id']}/stats", headers=h).json()["by_max_severity"]

    pdf = c.get(f"/v1/audits/{a['id']}/report.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    x = c.get(f"/v1/audits/{a['id']}/report.xlsx", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    xl = c.get(f"/v1/audits/{a['id']}/annotated.xlf", headers=h)
    assert xl.status_code == 200 and "ui.audited.xlf" in xl.headers["content-disposition"]
    root = etree.fromstring(xl.content)
    ns = {"x": "urn:oasis:names:tc:xliff:document:1.2"}
    notes = root.findall(".//x:trans-unit[@id='4']/x:note", ns)
    assert notes and notes[0].get("from") == "arbiter-auditor" and "tag_missing" in notes[0].text
    # the annotated file is still a valid bilingual XLIFF: readers see the same segments
    assert len(bilingual.parse_xliff(xl.content)) == 5
    lst = c.get("/v1/audits", headers=h).json()
    assert [i["id"] for i in lst["items"]] == [a["id"]]


def test_audit_xliff2_sample_mode_csv_tmx_and_pair(db):
    c, h = _client(db)
    r = c.post(
        "/v1/audits",
        headers=h,
        files={"file": ("doc.xliff", XLF2)},
        data={"source_lang": "en", "target_lang": "de", "ai_mode": "sample", "ai_sample_pct": "0.5"},
    ).json()
    run_until_idle()
    a = c.get(f"/v1/audits/{r['id']}", headers=h).json()
    assert a["state"] == "done" and a["summary"]["issues_by_code"].get("url_email_mismatch") == 1
    assert a["summary"]["ai"]["checked"] <= 1
    xl = c.get(f"/v1/audits/{a['id']}/annotated.xlf", headers=h)
    root = etree.fromstring(xl.content)
    ns = {"x": "urn:oasis:names:tc:xliff:document:2.0"}
    assert root.find(".//x:unit/x:notes/x:note", ns) is not None
    assert etree.QName(list(root.find(".//x:unit", ns))[0]).localname == "notes"

    for name, data in (("pairs.csv", CSV), ("mem.tmx", TMX)):
        r = c.post(
            "/v1/audits",
            headers=h,
            files={"file": (name, data)},
            data={"source_lang": "en", "target_lang": "de"},
        )
        assert r.status_code == 201, r.text
    run_until_idle()
    audits = {a["input_format"]: a for a in c.get("/v1/audits", headers=h).json()["items"]}
    assert (
        audits["csv"]["state"] == "done"
        and audits["csv"]["summary"]["issues_by_code"]["number_mismatch"] == 1
    )
    assert audits["tmx"]["state"] == "done" and audits["tmx"]["segments_total"] == 2
    x = c.get(f"/v1/audits/{audits['csv']['id']}/report.xlsx", headers=h).content
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(x))["Errors"]
    assert all(not str(c.value or "").startswith("=") for row in ws.iter_rows() for c in row)
    fresh = c.get(f"/v1/audits/{audits['tmx']['id']}/annotated.xlf", headers=h)
    assert b'version="2.1"' in fresh.content

    src = b"# Terms\n\nThe service starts on 1 March. Payment is due in 30 days.\n"
    tgt = b"# Uslovi\n\nUsluga pocinje 1. marta. Placanje dospeva za 31 dan.\n"
    r = c.post(
        "/v1/audits",
        headers=h,
        files={"file": ("en.md", src), "target_file": ("sr.md", tgt)},
        data={"source_lang": "en", "target_lang": "sr", "threshold": "50"},
    ).json()
    assert r["input_format"] == "pair"
    run_until_idle()
    a = c.get(f"/v1/audits/{r['id']}", headers=h).json()
    assert a["state"] == "done" and a["segments_total"] == 3
    assert a["summary"]["issues_by_code"]["number_mismatch"] == 1
    assert a["summary"]["scorecard"]["threshold"] == 50


def test_audit_failure_retry_validation_and_tenancy(db):
    c, h = _client(db)
    assert (
        c.post(
            "/v1/audits",
            headers=h,
            files={"file": ("a.docx", b"x")},
            data={"source_lang": "en", "target_lang": "de"},
        ).status_code
        == 422
    )
    r = c.post(
        "/v1/audits",
        headers=h,
        files={
            "file": ("bad.xlf", b"<xliff"),
        },
        data={"source_lang": "en", "target_lang": "de"},
    ).json()
    run_until_idle()
    a = c.get(f"/v1/audits/{r['id']}", headers=h).json()
    assert a["state"] == "failed" and "well-formed" in a["error"]
    assert c.get(f"/v1/audits/{a['id']}/report.pdf", headers=h).status_code == 409
    r = c.post(f"/v1/audits/{a['id']}/retry", headers=h)
    assert r.status_code == 200 and r.json()["state"] == "queued"
    assert db.execute(select(WorkItem).where(WorkItem.kind == "audit.parse")).scalars().all()
    run_until_idle()
    assert c.get(f"/v1/audits/{a['id']}", headers=h).json()["state"] == "failed"
    assert (
        c.post(
            "/v1/audits",
            headers=h,
            files={"file": ("ui.xlf", XLF12)},
            data={"source_lang": "en", "target_lang": "de", "glossary_id": "gls_nope"},
        ).status_code
        == 404
    )
    other = TestClient(create_app())
    r2 = other.post(
        "/v1/auth/register",
        json={"org_name": "Other", "name": "O", "email": "o@example.com", "password": "secret-123"},
    )
    h2 = {"Authorization": f"Bearer {r2.json()['token']}"}
    assert c.get(f"/v1/audits/{a['id']}", headers=h2).status_code == 404
    org = db.execute(select(Organization).where(Organization.name == "Acme")).scalar_one()
    org.settings = {**(org.settings or {}), "modules": ["tms"]}
    db.commit()
    r = c.get("/v1/audits", headers=h)
    assert r.status_code == 404 and r.json()["error"]["code"] == "module_disabled"
    assert db.execute(select(Audit)).scalars().first() is not None
