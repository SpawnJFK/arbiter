from __future__ import annotations

import copy
import io
import zipfile

import docx
import pytest
from fileproc_fixtures import TOC_AND_TEXTBOX, W, joined, make_docx, plain, tagged
from lxml import etree
from roundtrip import identity, translation

from arbiter.fileproc.base import FormatError, InlineCode, source_targets
from arbiter.fileproc.docx import DocxHandler

R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
TRANSLATED_PARTS = ("word/document.xml", "word/header1.xml", "word/footer1.xml", "word/footnotes.xml")


def _canon(rpr: etree._Element | None) -> str:
    if rpr is None:
        return ""
    c = copy.deepcopy(rpr)
    for el in c.iter():
        for a in list(el.attrib):
            if a.rsplit("}", 1)[-1].startswith("rsid"):
                del el.attrib[a]
    for child in list(c):
        if child.tag == f"{{{W}}}lang":
            c.remove(child)
    return etree.tostring(c, method="c14n").decode() if len(c) or c.attrib else ""


def signature(data: bytes) -> dict[str, list[list[tuple]]]:
    """Per part, per paragraph: visible text runs with effective formatting and link, plus markup items."""
    out: dict[str, list[list[tuple]]] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for part in TRANSLATED_PARTS:
            root = etree.fromstring(z.read(part))
            paras = []
            for p in root.iter(f"{{{W}}}p"):
                seq: list[tuple] = []
                for el in p.iter():
                    if not isinstance(el.tag, str) or el is p:
                        continue
                    name = el.tag.rsplit("}", 1)[-1]
                    run = next((a for a in el.iterancestors() if a.tag == f"{{{W}}}r"), None)
                    key = _canon(run.find(f"{{{W}}}rPr")) if run is not None else ""
                    link = next(
                        (a.get(R_ID) for a in el.iterancestors() if a.tag == f"{{{W}}}hyperlink"), None
                    )
                    if name == "t":
                        if seq and seq[-1][0] == "t" and seq[-1][2:] == (key, link):
                            seq[-1] = ("t", seq[-1][1] + (el.text or ""), key, link)
                        else:
                            seq.append(("t", el.text or "", key, link))
                    elif name in (
                        "tab",
                        "br",
                        "drawing",
                        "footnoteReference",
                        "footnoteRef",
                        "fldChar",
                        "instrText",
                        "bookmarkStart",
                        "bookmarkEnd",
                        "fldSimple",
                    ):
                        seq.append((name, key, etree.tostring(el) if name == "instrText" else b""))
                paras.append(seq)
            out[part] = paras
    return out


def parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_extraction_model() -> None:
    result = DocxHandler().extract(make_docx(), "en")
    units = {uid: (segs, ctx) for uid, segs, ctx in tagged(result)}
    # Split runs with identical formatting are merged, proofErr and w:lang differences vanish.
    assert units["word/document.xml#0"] == (
        ["This is ⟦1⟧important⟦/1⟧ and ⟦2⟧styled⟦/2⟧ text in one sentence.", "It has a second sentence."],
        "body",
    )
    assert units["word/document.xml#1"][0] == ["Visit ⟦1⟧our website⟦/1⟧ for more details."]
    assert units["word/document.xml#2"][0] == ["Name:⟦1/⟧John Smith⟦2/⟧"]
    # The whole field (instruction and result) is one standalone code.
    assert units["word/document.xml#3"][0] == ["Page ⟦1/⟧ of the annual report"]
    assert units["word/document.xml#4"][0] == ["Logo here: ⟦1/⟧ and the text goes on."]
    assert units["word/document.xml#5"][0] == ["⟦1/⟧Bookmarked text⟦2/⟧"]
    assert units["word/document.xml#6"][0] == ["Line one⟦1/⟧Line two"]
    contexts = {ctx for _, _, ctx in tagged(result)}
    assert {"body", "table", "header", "footer", "footnote"} <= contexts
    texts = [plain(joined(u)) for u in result.units]
    assert "First cell" in texts and "Second cell" in texts and "Nested cell" in texts
    assert "12345" not in texts and "Reviewer comment, not translated" not in texts
    assert any("This is a footnote." in t for t in texts)
    assert "Company header" in texts

    codes = {
        c.display for u in result.units for s in u.segments for c in s.content if isinstance(c, InlineCode)
    }
    assert {
        "bold",
        "italic",
        "link",
        "tab",
        "footnote",
        "field PAGE",
        "image",
        "bookmark",
        "line break",
        "footnote number",
    } <= codes


def test_identity_is_byte_identical() -> None:
    data = make_docx()
    merged, _ = identity(DocxHandler(), data)
    assert merged == data


def test_identity_rebuild_preserves_text_and_formatting() -> None:
    data = make_docx()
    merged, _ = identity(DocxHandler(rewrite_unchanged=True), data)
    before, after = parts(data), parts(merged)
    assert list(before) == list(after)
    for name in before:
        if name not in TRANSLATED_PARTS:
            assert before[name] == after[name], name
    assert signature(merged) == signature(data)
    reopened = docx.Document(io.BytesIO(merged))
    assert [p.text for p in reopened.paragraphs] == [
        p.text for p in docx.Document(io.BytesIO(data)).paragraphs
    ]


def test_translation_roundtrip_keeps_codes_and_formatting() -> None:
    data = make_docx()
    merged, _, _ = translation(DocxHandler(), data)
    sig = signature(merged)["word/document.xml"]
    flat = [item for para in sig for item in para]
    bold = [t for t in flat if t[0] == "t" and "IMPORTANT" in t[1]]
    assert bold and "<w:b" in bold[0][2]
    italic = [t for t in flat if t[0] == "t" and "STYLED" in t[1]]
    assert italic and "<w:i" in italic[0][2]
    link = [t for t in flat if t[0] == "t" and "OUR WEBSITE" in t[1]]
    assert link and link[0][3] == "rId10"
    kinds = [t[0] for t in flat]
    for needed in ("tab", "footnoteReference", "fldChar", "instrText", "drawing", "bookmarkStart", "br"):
        assert needed in kinds, needed
    # The field instruction itself is never translated.
    assert any(t[0] == "instrText" and b" PAGE " in t[2] for t in flat)
    # Leading space before the bold word survives with xml:space.
    root = etree.fromstring(parts(merged)["word/document.xml"])
    for t in root.iter(f"{{{W}}}t"):
        if t.text and t.text != t.text.strip():
            assert t.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    footnotes = parts(merged)["word/footnotes.xml"].decode()
    assert "THIS IS A FOOTNOTE.]]" in footnotes
    assert "[[COMPANY HEADER]]" in parts(merged)["word/header1.xml"].decode()
    assert parts(merged)["word/comments.xml"] == parts(data)["word/comments.xml"]
    docx.Document(io.BytesIO(merged))


def test_tracked_changes_rejected() -> None:
    with pytest.raises(FormatError, match="tracked changes"):
        DocxHandler().extract(make_docx(tracked=True), "en")


def test_standalone_codes_can_be_reordered() -> None:
    data = make_docx()
    h = DocxHandler()
    result = h.extract(data, "en")
    unit = next(u for u in result.units if u.unit_id == "word/document.xml#2")
    seg = unit.segments[0].content  # Name:, tab, John Smith, footnote
    tab, footnote = seg[1], seg[3]
    targets = source_targets(result)
    targets[unit.unit_id] = [["Ime", footnote, " i prezime:", tab, "Jovan"]]
    merged = h.merge(data, targets, "sr")
    again = {u.unit_id: u for u in h.extract(merged, "en").units}
    content = joined(again[unit.unit_id])
    displays = [c.display for c in content if isinstance(c, InlineCode)]
    assert displays == ["footnote", "tab"]
    assert plain(content) == "Ime i prezime:Jovan"


def test_unknown_code_rejected() -> None:
    data = make_docx()
    h = DocxHandler()
    result = h.extract(data, "en")
    targets = source_targets(result)
    targets["word/document.xml#2"] = [["Name", InlineCode("99", "standalone")]]
    with pytest.raises(FormatError, match="99"):
        h.merge(data, targets, "de")


def test_broken_pairs_and_dropped_standalone_rejected() -> None:
    data = make_docx()
    h = DocxHandler()
    result = h.extract(data, "en")
    unit = result.units[0]
    seg = unit.segments[0].content
    opener = next(c for c in seg if isinstance(c, InlineCode) and c.kind == "open")
    bad = source_targets(result)
    bad[unit.unit_id] = [["Only ", opener, "open"], unit.segments[1].content]
    with pytest.raises(FormatError):
        h.merge(data, bad, "de")
    tab_unit = next(u for u in result.units if u.unit_id == "word/document.xml#2")
    dropped = source_targets(result)
    dropped[tab_unit.unit_id] = [["No codes at all"]]
    with pytest.raises(FormatError, match="missing"):
        h.merge(data, dropped, "de")


def test_damaged_file() -> None:
    with pytest.raises(FormatError, match="damaged"):
        DocxHandler().extract(b"PK\x03\x04 not really a zip", "en")


def test_field_spanning_paragraphs_and_text_boxes() -> None:
    data = make_docx(extra=TOC_AND_TEXTBOX)
    h = DocxHandler()
    result = h.extract(data, "en")
    texts = [plain(joined(u)) for u in result.units]
    # TOC entries are field results: not translatable, the paragraph holding only them is skipped.
    assert not any("Introduction chapter" in t for t in texts)
    assert "Contents " in texts or "Contents" in [t.strip() for t in texts]
    assert "After the table of contents." in texts
    assert texts.count("Text box words") == 1
    assert any("legacy copy" in w for w in result.warnings)
    merged, _, _ = translation(h, data)
    root = etree.fromstring(parts(merged)["word/document.xml"])
    kinds = [fc.get(f"{{{W}}}fldCharType") for fc in root.iter(f"{{{W}}}fldChar")]
    assert kinds.count("begin") == kinds.count("end")
    xml = parts(merged)["word/document.xml"].decode()
    assert "TEXT BOX WORDS" in xml and "Introduction chapter" in xml and "PAGEREF _Toc1" in xml
    merged2, _ = identity(DocxHandler(rewrite_unchanged=True), data)
    assert signature(merged2) == signature(data)
