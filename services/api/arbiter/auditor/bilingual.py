"""Readers for bilingual files produced by any tool: XLIFF 1.2 and 2.x, TMX, bilingual
CSV and XLSX (source and target columns), or a source + target document pair aligned
sentence by sentence (linguistic/align.py).

Every reader returns BiSegment rows in tagged text (⟦n⟧ codes, ⟦⟦ ⟧⟧ escapes). In XLIFF the
inline elements of source and target share ids, so their codes line up and the tag
checks can compare them. Every failure is a FormatError the client can read.
"""

from __future__ import annotations

import csv
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from typing import Any

from lxml import etree

from arbiter.fileproc.base import LB, RB, FormatError, to_tagged
from arbiter.fileproc.registry import detect_and_extract
from arbiter.linguistic import tm
from arbiter.linguistic.align import align

X12 = "urn:oasis:names:tc:xliff:document:1.2"
X2 = "urn:oasis:names:tc:xliff:document:2.0"
FORMATS = ("xliff", "tmx", "csv", "xlsx")
MAX_SEGMENTS = 200_000
_BAD_ID = re.compile(r"[^A-Za-z0-9_.-]")


@dataclass
class BiSegment:
    ref: str  # "<file index>|<unit id>|<segment id>" for XLIFF, row or tuid otherwise
    source_tagged: str
    target_tagged: str


def _esc(text: str) -> str:
    return text.replace(LB, LB + LB).replace(RB, RB + RB)


def detect_format(filename: str) -> str | None:
    ext = (
        posixpath.basename(filename.replace("\\", "/")).rsplit(".", 1)[-1].lower() if "." in filename else ""
    )
    return {"xlf": "xliff", "xliff": "xliff", "sdlxliff": "xliff", "mqxliff": "xliff", "tmx": "tmx",
            "csv": "csv", "tsv": "csv", "xlsx": "xlsx"}.get(ext)  # fmt: skip


def _parser() -> etree.XMLParser:
    return etree.XMLParser(resolve_entities=False, no_network=True, remove_comments=True, huge_tree=True)


def _local(el: Any) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def _cid(value: str | None) -> str:
    return _BAD_ID.sub("_", value or "x") or "x"


def xliff_tagged(el: Any, version: int) -> str:
    """Inline content of an XLIFF source/target/mrk element as tagged text."""
    out: list[str] = []

    def walk(parent: Any) -> None:
        if parent.text:
            out.append(_esc(parent.text))
        for child in parent:
            name = _local(child)
            if version == 2:
                if name == "pc":
                    cid = _cid(child.get("id"))
                    out.append(f"{LB}{cid}{RB}")
                    walk(child)
                    out.append(f"{LB}/{cid}{RB}")
                elif name == "ph":
                    out.append(f"{LB}{_cid(child.get('id'))}/{RB}")
                elif name == "sc":
                    out.append(f"{LB}{_cid(child.get('id'))}{RB}")
                elif name == "ec":
                    out.append(f"{LB}/{_cid(child.get('startRef') or child.get('id'))}{RB}")
                elif name == "cp":
                    out.append(_esc(chr(int(child.get("hex", "20"), 16))))
                elif name:
                    walk(child)
            else:
                if name == "g":
                    cid = _cid(child.get("id"))
                    out.append(f"{LB}{cid}{RB}")
                    walk(child)
                    out.append(f"{LB}/{cid}{RB}")
                elif name in ("x", "ph", "it"):
                    out.append(f"{LB}{_cid(child.get('id'))}/{RB}")
                elif name in ("bx", "bpt"):
                    out.append(f"{LB}{_cid(child.get('rid') or child.get('id'))}{RB}")
                elif name in ("ex", "ept"):
                    out.append(f"{LB}/{_cid(child.get('rid') or child.get('id'))}{RB}")
                elif name and name != "sub":
                    walk(child)
            if child.tail:
                out.append(_esc(child.tail))

    walk(el)
    return "".join(out)


def iter_xliff(root: Any) -> list[tuple[Any, str, Any, Any, int]]:
    """(unit element, ref, source element, target element or None, version) per segment."""
    ns = etree.QName(root).namespace
    out: list[tuple[Any, str, Any, Any, int]] = []
    if ns == X12:
        for fi, f in enumerate(root.iter(f"{{{X12}}}file")):
            for tu in f.iter(f"{{{X12}}}trans-unit"):
                if tu.get("translate") == "no":
                    continue
                uid = tu.get("id", "")
                seg_source = tu.find(f"{{{X12}}}seg-source")
                target = tu.find(f"{{{X12}}}target")
                mrks = (
                    [m for m in seg_source.iter(f"{{{X12}}}mrk") if m.get("mtype") == "seg"]
                    if seg_source is not None
                    else []
                )
                if mrks:
                    tmrks = {}
                    if target is not None:
                        tmrks = {
                            m.get("mid"): m for m in target.iter(f"{{{X12}}}mrk") if m.get("mtype") == "seg"
                        }
                    for m in mrks:
                        out.append((tu, f"{fi + 1}|{uid}|{m.get('mid', '')}", m, tmrks.get(m.get("mid")), 1))
                    continue
                source = tu.find(f"{{{X12}}}source")
                if source is not None:
                    out.append((tu, f"{fi + 1}|{uid}|", source, target, 1))
        return out
    if ns == X2:
        for fi, f in enumerate(root.iter(f"{{{X2}}}file")):
            for u in f.iter(f"{{{X2}}}unit"):
                if u.get("translate") == "no":
                    continue
                for si, seg in enumerate(u.iter(f"{{{X2}}}segment")):
                    src = seg.find(f"{{{X2}}}source")
                    if src is None:
                        continue
                    sid = seg.get("id") or str(si + 1)
                    out.append((u, f"{fi + 1}|{u.get('id', '')}|{sid}", src, seg.find(f"{{{X2}}}target"), 2))
        return out
    raise FormatError("The file is not a supported XLIFF version (1.2, 2.0 or 2.1).")


def parse_xliff(data: bytes) -> list[BiSegment]:
    try:
        root = etree.fromstring(data, _parser())
    except etree.XMLSyntaxError:
        raise FormatError("The XLIFF file is not well-formed XML.") from None
    out = []
    for _unit, ref, src, tgt, version in iter_xliff(root):
        source = xliff_tagged(src, version)
        if not tm.tagged_plain(source).strip():
            continue
        out.append(BiSegment(ref, source, xliff_tagged(tgt, version) if tgt is not None else ""))
    return out


def parse_tmx(data: bytes, source_lang: str, target_lang: str) -> list[BiSegment]:
    try:
        rows = tm.tmx_pairs(data, source_lang, target_lang)
    except etree.XMLSyntaxError:
        raise FormatError("The TMX file is not well-formed XML.") from None
    return [BiSegment(f"tu:{tuid}", s, t) for tuid, s, t in rows]


_SRC_HEAD = re.compile(r"^(source|src|original|source text|source_text)\b", re.I)
_TGT_HEAD = re.compile(r"^(target|tgt|translation|target text|target_text)\b", re.I)


def _columns(rows: list[list[str]]) -> tuple[int, int, int]:
    """(source column, target column, first data row)."""
    if rows:
        head = [str(c or "").strip() for c in rows[0]]
        s = next((i for i, h in enumerate(head) if _SRC_HEAD.match(h)), None)
        t = next((i for i, h in enumerate(head) if _TGT_HEAD.match(h)), None)
        if s is not None and t is not None:
            return s, t, 1
    return 0, 1, 0


def _from_rows(rows: list[list[Any]], prefix: str) -> list[BiSegment]:
    s, t, start = _columns([[str(c) if c is not None else "" for c in r] for r in rows[:1]])
    out = []
    for i, row in enumerate(rows[start:], start=start + 1):
        src = str(row[s]) if len(row) > s and row[s] is not None else ""
        tgt = str(row[t]) if len(row) > t and row[t] is not None else ""
        if src.strip():
            out.append(BiSegment(f"{prefix}{i}", _esc(src), _esc(tgt)))
    return out


def parse_csv(data: bytes) -> list[BiSegment]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise FormatError("The CSV file must be UTF-8 encoded.") from None
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows or max(len(r) for r in rows) < 2:
        raise FormatError("A bilingual CSV needs a source and a target column.")
    return _from_rows(rows, "row:")


def parse_xlsx(data: bytes) -> list[BiSegment]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise FormatError("The XLSX file could not be read.") from None
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows or max(len(r) for r in rows) < 2:
        raise FormatError("A bilingual XLSX needs a source and a target column in its first sheet.")
    return _from_rows(rows, "row:")


def sentences(filename: str, data: bytes, lang: str) -> list[str]:
    result = detect_and_extract(filename, data, lang)
    return [to_tagged(s.content) for u in result.units for s in u.segments]


def parse_pair(
    src_name: str, src_data: bytes, tgt_name: str, tgt_data: bytes, source_lang: str, target_lang: str
) -> tuple[list[BiSegment], int]:
    """Align a source document with its translation. Returns (segments, unaligned target sentences).
    Codes of two separately extracted files do not share ids, so both sides are plain text."""
    src = [tm.tagged_plain(s) for s in sentences(src_name, src_data, source_lang)]
    tgt = [tm.tagged_plain(t) for t in sentences(tgt_name, tgt_data, target_lang)]
    out: list[BiSegment] = []
    extra = 0
    for n, bead in enumerate(align(src, tgt), start=1):
        if not bead.src:
            extra += len(bead.tgt)
            continue
        out.append(
            BiSegment(
                f"pair:{n}",
                _esc(" ".join(src[i] for i in bead.src)),
                _esc(" ".join(tgt[j] for j in bead.tgt)),
            )
        )
    return out, extra


def read(
    fmt: str,
    data: bytes,
    source_lang: str,
    target_lang: str,
    *,
    filename: str = "",
    target_name: str | None = None,
    target_data: bytes | None = None,
) -> tuple[list[BiSegment], dict[str, Any]]:
    """Dispatch on the input format. Returns (segments, notes for the summary)."""
    notes: dict[str, Any] = {}
    if fmt == "pair":
        assert target_name is not None and target_data is not None
        segs, extra = parse_pair(filename, data, target_name, target_data, source_lang, target_lang)
        notes["unaligned_target_sentences"] = extra
    elif fmt == "xliff":
        segs = parse_xliff(data)
    elif fmt == "tmx":
        segs = parse_tmx(data, source_lang, target_lang)
    elif fmt == "csv":
        segs = parse_csv(data)
    elif fmt == "xlsx":
        segs = parse_xlsx(data)
    else:
        raise FormatError("Supported audit inputs: XLIFF, TMX, bilingual CSV or XLSX, or a document pair.")
    if not segs:
        raise FormatError("The file contains no bilingual segments to audit.")
    if len(segs) > MAX_SEGMENTS:
        raise FormatError(f"The file has more than {MAX_SEGMENTS} segments; split it into smaller files.")
    return segs, notes
